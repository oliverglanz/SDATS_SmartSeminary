from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from app.banner_donsheet_comparison import BANNER_RENAMES, SITE_NAMES
from app.config import PROJECT_ROOT


TEMPLATE_DIRECTORY = PROJECT_ROOT / "0_source_files" / "default_DonSheet"
MAX_DATA_ROWS = 300

SOURCE_TO_TEMPLATE_HEADER = {
    "Course Start Date {Meet Start Date}": "Intensive Period Start Date {Meet Start Date}",
    "Course End Date {Meet End Date}": "Intensive Period End Date {Meet End Date}",
}


@dataclass(frozen=True)
class PrefilledDonSheetResult:
    excel_bytes: bytes
    output_filename: str
    rows_by_department: dict[str, int]
    template_filename: str


def _unique_headers(headers) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, value in enumerate(headers, start=1):
        name = "" if value is None else str(value).strip()
        name = name or f"Unnamed_{index}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        result.append(name if count == 0 else f"{name}_{count}")
    return result


def _load_banner(data: bytes, filename: str) -> pd.DataFrame:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else "xlsx"
    if extension == "csv":
        frame = pd.read_csv(BytesIO(data), dtype=str, na_filter=False)
    else:
        raw = pd.read_excel(BytesIO(data), header=None, nrows=25, dtype=str, na_filter=False)
        header_row = next(
            (index for index, row in raw.iterrows() if "CRN" in {str(value).strip() for value in row}),
            0,
        )
        frame = pd.read_excel(BytesIO(data), header=header_row, dtype=str, na_filter=False)

    frame.columns = _unique_headers(frame.columns)
    frame = frame.rename(columns=BANNER_RENAMES)
    if "CRN" not in frame.columns:
        raise ValueError("The Banner Sheet must contain a CRN column.")

    for column in frame.columns:
        values = frame[column].astype("string").str.strip()
        frame[column] = values.mask(values.str.lower().isin({"", "nan", "none", "na", "n/a"}))

    if "Building {Meet Bldg}" in frame.columns:
        frame["Building {Meet Bldg}"] = frame["Building {Meet Bldg}"].replace(SITE_NAMES)

    frame["CRN"] = frame["CRN"].astype("string").str.strip()
    frame = frame[frame["CRN"].notna() & frame["CRN"].ne("")].copy()
    frame["CRN sorted {not in Banner}"] = (
        frame["CRN"] + "-" + frame.groupby("CRN", dropna=False).cumcount().add(1).astype(str).str.zfill(2)
    )
    return frame


def _filter_department_rows(banner: pd.DataFrame) -> dict[str, pd.DataFrame]:
    required = ["Subject", "Campus {Ssasect Campus}", "Course Number {Crse Num}"]
    missing = [column for column in required if column not in banner.columns]
    if missing:
        raise ValueError(f"The Banner Sheet is missing required column(s): {', '.join(missing)}.")

    working = banner.copy()
    working["Subject"] = working["Subject"].astype("string").str.strip()
    working["Campus {Ssasect Campus}"] = working["Campus {Ssasect Campus}"].astype("string").str.strip()
    working["Course Number {Crse Num}"] = pd.to_numeric(
        working["Course Number {Crse Num}"], errors="coerce"
    )

    if "Scacrse College" in working.columns:
        working = working[working["Scacrse College"].astype("string").str.strip().eq("70")].copy()
    if "Status" in working.columns:
        working = working[~working["Status"].astype("string").str.strip().str.upper().eq("I")].copy()
    restriction_columns = [column for column in ["Lvl Res Ind", "Cmp Res Ind"] if column in working.columns]
    if restriction_columns:
        working = working[working[restriction_columns].isna().all(axis=1)].copy()

    main_campus = working["Campus {Ssasect Campus}"].isin(["0", "7"])
    outside_dmin_range = (
        working["Course Number {Crse Num}"].lt(700) | working["Course Number {Crse Num}"].gt(799)
    )
    standard_departments = ["CHIS", "DSLE", "GSEM", "MSSN", "NTST", "PATH", "THST"]
    routed = {
        department: working.loc[
            working["Subject"].eq(department) & outside_dmin_range & main_campus
        ].copy()
        for department in standard_departments
    }
    routed["OTST"] = working.loc[
        working["Subject"].isin(["OTST", "ANEA"]) & outside_dmin_range & main_campus
    ].copy()
    routed["MDivHISP_MAPmENGL_MAPmHISP"] = working.loc[
        working["Campus {Ssasect Campus}"].eq("9OM")
    ].copy()
    routed["MA_Religion"] = working.loc[
        working["Campus {Ssasect Campus}"].eq("ZOM")
    ].copy()
    routed["DMIN"] = working.loc[
        working["Course Number {Crse Num}"].between(700, 799, inclusive="both")
    ].copy()
    return routed


def _latest_formula_template() -> Path:
    candidates = [
        path
        for path in TEMPLATE_DIRECTORY.glob("*.xlsx")
        if "withformulas" in path.name.lower() or "withfomulas" in path.name.lower()
    ]
    if not candidates:
        raise FileNotFoundError(
            "No formula-enabled DonSheet template was found in "
            f"{TEMPLATE_DIRECTORY}."
        )
    return max(candidates, key=lambda path: path.stat().st_mtime)


def _output_filename(banner_filename: str) -> str:
    name = Path(banner_filename).stem.upper()
    year = re.search(r"(?<!\d)(20\d{2})(?!\d)", name)
    semester = re.search(r"(SPRING|SUMMER|FALL)", name)
    if year and semester:
        return f"DonSheet_{year.group(1)}-{semester.group(1).title()}.xlsx"
    return f"DonSheet_prefilled_{Path(banner_filename).stem}.xlsx"


def _is_formula(cell) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def prefill_donsheet_from_banner(banner_bytes: bytes, banner_filename: str) -> PrefilledDonSheetResult:
    """Populate the latest formula-enabled DonSheet template from one Banner export."""
    departments = _filter_department_rows(_load_banner(banner_bytes, banner_filename))
    template_path = _latest_formula_template()
    workbook = load_workbook(BytesIO(template_path.read_bytes()), data_only=False, keep_links=True)

    for sheet_name, dataframe in departments.items():
        if sheet_name not in workbook.sheetnames:
            raise KeyError(f"The DonSheet template is missing expected worksheet '{sheet_name}'.")
        if len(dataframe) > MAX_DATA_ROWS - 1:
            raise ValueError(
                f"{sheet_name} has {len(dataframe)} Banner rows, exceeding the template capacity of "
                f"{MAX_DATA_ROWS - 1}."
            )

        worksheet = workbook[sheet_name]
        headers = {str(cell.value).strip(): cell.column for cell in worksheet[1] if cell.value is not None}
        for row in worksheet.iter_rows(min_row=2, max_row=MAX_DATA_ROWS, max_col=worksheet.max_column):
            for cell in row:
                if not _is_formula(cell):
                    cell.value = None

        for source_header in dataframe.columns:
            target_header = SOURCE_TO_TEMPLATE_HEADER.get(source_header, source_header)
            target_column = headers.get(target_header)
            if target_column is None:
                continue
            for row_number, value in enumerate(dataframe[source_header], start=2):
                destination = worksheet.cell(row=row_number, column=target_column)
                if not _is_formula(destination):
                    destination.value = None if pd.isna(value) else value

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except AttributeError:
        pass

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return PrefilledDonSheetResult(
        excel_bytes=stream.getvalue(),
        output_filename=_output_filename(banner_filename),
        rows_by_department={name: len(frame) for name, frame in departments.items()},
        template_filename=template_path.name,
    )
